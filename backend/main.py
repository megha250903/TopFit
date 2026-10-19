import io
import json
import os
from typing import Optional, List
from dotenv import load_dotenv
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, status
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from groq import Groq
from pypdf import PdfReader
from docx import Document

load_dotenv()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

if not GROQ_API_KEY:
    print("Warning: GROQ_API_KEY not found in environment!")

client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None
PRIMARY_MODEL = "openai/gpt-oss-120b"
FALLBACK_MODEL = "openai/gpt-oss-20b"

app = FastAPI(
    title="TopFit AI Resume Screening API",
    description="FastAPI backend for job parsing and resume scoring",
    version="1.0.0",
)

# CORS configuration
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173", "*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Pydantic schemas
class JobResponse(BaseModel):
    role: str = "Candidate Role"
    required_skills: List[str] = []
    preferred_skills: List[str] = []
    minimum_experience: Optional[float] = None
    education_requirements: List[str] = []
    responsibilities: List[str] = []

class ScoreResponse(BaseModel):
    filename: str
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    experience_years: Optional[float] = None
    score: int
    matching_skills: List[str] = []
    missing_skills: List[str] = []
    experience_met: Optional[bool] = None
    verdict: str

def call_groq_json(system_prompt: str, user_prompt: str) -> dict:
    if not client:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Groq API client is not configured. Please set GROQ_API_KEY in backend/.env.",
        )
    
    # Try primary model first, fallback if unavailable
    for target_model in [PRIMARY_MODEL, FALLBACK_MODEL]:
        try:
            response = client.chat.completions.create(
                model=target_model,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                response_format={"type": "json_object"},
                temperature=0.2,
            )
            content = response.choices[0].message.content
            return json.loads(content)
        except Exception as e:
            print(f"Model {target_model} failed: {e}")
            if target_model == FALLBACK_MODEL:
                raise HTTPException(
                    status_code=status.HTTP_502_BAD_GATEWAY,
                    detail=f"AI model error: {str(e)}",
                )
    return {}

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    text = ""
    try:
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        elif ext == ".docx":
            doc = Document(io.BytesIO(file_bytes))
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + "\n"
        elif ext == ".txt":
            text = file_bytes.decode("utf-8", errors="ignore")
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file format '{ext}'. Please upload a PDF, DOCX, or TXT file.",
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read {filename}. The file may be corrupted or encrypted: {str(e)}",
        )
    return text.strip()

@app.get("/api/health")
def health():
    return {"status": "ok", "service": "topfit-backend"}

@app.post("/api/job", response_model=JobResponse)
async def parse_job(
    text: Optional[str] = Form(None),
    file: Optional[UploadFile] = File(None),
):
    job_content = ""
    if file:
        content_bytes = await file.read()
        job_content = extract_text_from_file(content_bytes, file.filename or "job.txt")
    elif text and text.strip():
        job_content = text.strip()

    if not job_content or len(job_content) < 30:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Job description must contain at least 30 characters of readable text.",
        )

    system_prompt = """
    You are an expert HR assistant. Analyze the provided job description and extract structured information.
    Return ONLY a JSON object with this exact structure:
    {
      "role": "Role Title",
      "required_skills": ["skill1", "skill2"],
      "preferred_skills": ["skill1", "skill2"],
      "minimum_experience": 2.0 (number of years, or null if not specified),
      "education_requirements": ["degree1"],
      "responsibilities": ["resp1", "resp2"]
    }
    If information is missing, use empty arrays [] or null. Do not invent information.
    """

    data = call_groq_json(system_prompt, f"Analyze this job description:\n\n{job_content}")

    # Sanitize and ensure types
    min_exp = data.get("minimum_experience")
    if min_exp is not None:
        try:
            min_exp = float(min_exp)
        except (ValueError, TypeError):
            min_exp = None

    return JobResponse(
        role=str(data.get("role") or "Target Position"),
        required_skills=list(data.get("required_skills") or []),
        preferred_skills=list(data.get("preferred_skills") or []),
        minimum_experience=min_exp,
        education_requirements=list(data.get("education_requirements") or []),
        responsibilities=list(data.get("responsibilities") or []),
    )

@app.post("/api/score", response_model=ScoreResponse)
async def score_resume(
    job: str = Form(...),
    file: UploadFile = File(...),
):
    try:
        job_dict = json.loads(job)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid job JSON format provided.",
        )

    filename = file.filename or "candidate_resume"
    file_bytes = await file.read()
    resume_text = extract_text_from_file(file_bytes, filename)

    if not resume_text or len(resume_text.strip()) < 20:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="This PDF has no readable text. It may be a scan. Try a DOCX or a text-based PDF.",
        )

    system_prompt = """
    You are a professional HR recruiter and talent evaluator.
    Compare the candidate's resume with the target job description.
    Return ONLY a JSON object matching this schema:
    {
      "name": "Candidate full name, or null",
      "email": "email address or null",
      "phone": "phone number or null",
      "experience_years": 3.5 (float estimate of total years experience, or null),
      "score": 85 (integer from 0 to 100 representing overall match),
      "matching_skills": ["skills that candidate possesses which match role"],
      "missing_skills": ["important role skills the candidate is missing"],
      "experience_met": true/false/null (true if candidate meets minimum experience requirement),
      "verdict": "2-3 concise sentences providing an objective recruiter summary and recommendation"
    }
    Rules:
    - score must be an integer between 0 and 100.
    - matching_skills and missing_skills should be lists of strings.
    - verdict should clearly explain the fit, strengths, and any gaps.
    """

    user_prompt = f"""
    TARGET JOB:
    {json.dumps(job_dict, indent=2)}

    CANDIDATE RESUME ({filename}):
    {resume_text}
    """

    data = call_groq_json(system_prompt, user_prompt)

    raw_score = data.get("score", 50)
    try:
        final_score = int(round(float(raw_score)))
        final_score = max(0, min(100, final_score))
    except (ValueError, TypeError):
        final_score = 50

    exp_years = data.get("experience_years")
    if exp_years is not None:
        try:
            exp_years = round(float(exp_years), 1)
        except (ValueError, TypeError):
            exp_years = None

    exp_met = data.get("experience_met")
    if exp_met not in (True, False, None):
        exp_met = None

    return ScoreResponse(
        filename=filename,
        name=data.get("name"),
        email=data.get("email"),
        phone=data.get("phone"),
        experience_years=exp_years,
        score=final_score,
        matching_skills=list(data.get("matching_skills") or []),
        missing_skills=list(data.get("missing_skills") or []),
        experience_met=exp_met,
        verdict=str(data.get("verdict") or "Candidate review generated."),
    )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
